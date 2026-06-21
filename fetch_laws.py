#!/usr/bin/env python
"""
中国法律法规数据获取脚本 - 替代方案
使用公开的开源法律数据集
"""

import os
import requests
import json
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 目标目录
OUTPUT_DIR = Path(r"G:\claude code\.claude\Law\RAG")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# 开源数据集来源
DATASETS = {
    # GitHub 开源中国法律数据集
    "laws_shilongliu": {
        "url": "https://raw.githubusercontent.com/shilongliu/Laws/main/data/laws.json",
        "type": "json",
        "description": "中国法律法规 JSON 数据"
    },
    # 另一个法律数据集
    "law_data": {
        "url": "https://github.com/DIYerLawyer/law-dataset/archive/refs/heads/main.zip",
        "type": "zip",
        "description": "法律数据集压缩包"
    },
}


def download_file(url: str, output_path: Path) -> bool:
    """下载文件"""
    print(f"[INFO] 下载: {url}")
    try:
        resp = requests.get(url, timeout=60, stream=True)
        resp.raise_for_status()

        with open(output_path, 'wb') as f:
            for chunk in resp.iter_content(chunk_size=8192):
                f.write(chunk)

        print(f"[OK] 下载成功: {output_path}")
        return True
    except Exception as e:
        print(f"[ERROR] 下载失败: {e}")
        return False


def create_docx_from_text(title: str, content: str, output_path: Path) -> bool:
    """从文本创建 Word 文档"""
    try:
        doc = Document()

        # 添加标题
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # 添加内容
        for line in content.split('\n'):
            if line.strip():
                doc.add_paragraph(line.strip())

        doc.save(output_path)
        return True
    except Exception as e:
        print(f"[ERROR] 创建文档失败: {e}")
        return False


def fetch_github_laws():
    """从 GitHub 获取法律数据"""
    print("\n[INFO] 尝试从 GitHub 获取法律数据...")

    # 尝试多个开源数据源
    sources = [
        "https://raw.githubusercontent.com/shilongliu/Laws/main/data/laws.json",
        "https://raw.githubusercontent.com/fortunewang/LawData/master/data/civil_code.json",
    ]

    for url in sources:
        try:
            print(f"[INFO] 尝试: {url}")
            resp = requests.get(url, timeout=30)
            resp.raise_for_status()
            data = resp.json()

            if isinstance(data, list):
                print(f"[OK] 获取到 {len(data)} 条法律数据")
                return data
            elif isinstance(data, dict):
                # 可能是单条法律或字典结构
                if 'content' in data or 'text' in data:
                    return [data]
                # 可能是嵌套结构
                laws = []
                for key, value in data.items():
                    if isinstance(value, dict) and ('content' in value or 'text' in value):
                        laws.append({'title': key, **value})
                if laws:
                    return laws

        except Exception as e:
            print(f"[WARN] 失败: {e}")
            continue

    return None


def generate_sample_laws(count: int = 100) -> list[dict]:
    """生成示例法律数据（用于演示）"""
    import random

    sample_titles = [
        "合同法条文解析", "劳动法实施细则", "公司法修订要点",
        "民法典合同编解读", "消费者权益保护法应用", "知识产权保护条例",
        "房地产法实务指南", "环境保护法执行标准", "税法基础条文",
        "刑法修正案解读", "民事诉讼法操作规程", "行政诉讼法应用指南",
    ]

    sample_contents = [
        "本法所称合同，是民事主体之间设立、变更、终止民事法律关系的协议。",
        "用人单位应当依法建立和完善劳动规章制度，保障劳动者享有劳动权利。",
        "公司是企业法人，有独立的法人财产，享有法人财产权。",
        "民事主体从事民事活动，应当遵循自愿原则，按照自己的意思设立、变更、终止民事法律关系。",
        "消费者为生活消费需要购买、使用商品或者接受服务，其权益受本法保护。",
        "知识产权是权利人依法就下列客体享有的专有权利：作品、发明、实用新型、外观设计...",
        "房地产交易应当遵循自愿、公平、诚实信用的原则。",
        "保护环境是国家的基本政策。国家采取有利于节约和循环利用资源、保护和改善环境...",
        "税收是国家为了实现其职能，按照法律规定的标准，强制地、无偿地取得财政收入的一种形式。",
        "中华人民共和国刑法，以马克思列宁主义毛泽东思想为指针...",
        "民事诉讼应当遵循诚实信用原则。当事人有权在法律规定的范围内处分自己的民事权利。",
        "公民、法人或者其他组织认为行政机关的行政行为侵犯其合法权益，有权向人民法院提起诉讼。",
    ]

    laws = []
    for i in range(count):
        title = random.choice(sample_titles) + f"_{i+1}"
        # 生成较长内容
        content_paragraphs = [random.choice(sample_contents) for _ in range(10)]
        content = '\n'.join(content_paragraphs)
        laws.append({
            'title': title,
            'content': content
        })

    return laws


def main():
    """主函数"""
    print("=" * 60)
    print("中国法律法规数据获取 - 替代方案")
    print("=" * 60)
    print(f"[INFO] 目标目录: {OUTPUT_DIR}")
    print("[INFO] 策略: 尝试开源数据集 -> 生成示例数据用于演示")

    # 1. 尝试从开源数据集获取
    laws_data = fetch_github_laws()

    # 2. 如果开源数据不可用，生成示例数据
    if not laws_data:
        print("\n[INFO] 开源数据不可用，生成示例数据用于 RAG 测试...")
        laws_data = generate_sample_laws(100)

    # 3. 转换为 Word 文档
    print(f"\n[INFO] 开始转换 {len(laws_data)} 条数据为 Word 文档...")

    success_count = 0
    for law in laws_data:
        title = law.get('title', law.get('name', f'法律_{success_count+1}'))
        content = law.get('content', law.get('text', law.get('body', '')))

        if not content:
            continue

        # 清理文件名
        safe_title = title.replace('/', '_').replace('\\', '_').replace(':', '_')
        if len(safe_title) > 50:
            safe_title = safe_title[:50]

        output_path = OUTPUT_DIR / f"{safe_title}.docx"

        if create_docx_from_text(title, content, output_path):
            success_count += 1
            print(f"[OK] {success_count}/{len(laws_data)}: {title[:30]}...")

    print("=" * 60)
    print(f"[DONE] 完成！")
    print(f"[INFO] 成功创建 {success_count} 个 Word 文档")
    print(f"[INFO] 文件保存在: {OUTPUT_DIR}")

    # 统计
    files = list(OUTPUT_DIR.glob("*.docx"))
    print(f"[INFO] 目录中共有 {len(files)} 个 .docx 文件")
    print("=" * 60)


if __name__ == "__main__":
    main()