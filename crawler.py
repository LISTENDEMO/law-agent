#!/usr/bin/env python
"""
国家法律法规数据库爬虫 - 修正版
爬取 https://flk.npc.gov.cn 的法律法规文档
"""

import os
import time
import requests
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 目标目录
OUTPUT_DIR = Path(r"G:\claude code\.claude\Law\RAG")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# API 端点
BASE_URL = "https://flk.npc.gov.cn"
SEARCH_API = f"{BASE_URL}/law-search/search/list"
DETAIL_API = f"{BASE_URL}/law-search/search/flfgDetails"

# 请求头
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/149.0.0.0 Safari/537.36",
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json;charset=UTF-8",
    "Referer": f"{BASE_URL}/search",
}

# 正确的法规分类代码（从 enumData API 获取）
CATEGORY_CODES = {
    "宪法": [100],
    "法律": [101, 102, 110, 120, 130, 140, 150, 155, 160, 170, 180, 190, 195, 200],
    "行政法规": [201, 210, 215],
    "监察法规": [220],
    "地方法规": [221, 222, 230, 260, 270, 290, 295, 300, 305, 310],
    "司法解释": [311, 320, 330, 340, 350],
}

# 时效性代码
SXX_VALID = [3]  # 有效状态


def search_laws(category_ids: list[int], page_num: int = 1, page_size: int = 50) -> dict:
    """搜索法律法规列表"""
    payload = {
        "searchRange": 1,
        "sxrq": [],
        "gbrq": [],
        "searchType": 2,
        "sxx": SXX_VALID,  # 只获取有效状态
        "gbrqYear": [],
        "flfgCodeId": category_ids,
        "zdjgCodeId": [],
        "searchContent": "",
        "orderByParam": {"order": "-1", "sort": ""},
        "pageNum": page_num,
        "pageSize": page_size,
    }

    try:
        resp = requests.post(SEARCH_API, headers=HEADERS, json=payload, timeout=30)
        resp.raise_for_status()
        data = resp.json()
        if data.get("code") == 200:
            return {"total": data.get("total", 0), "rows": data.get("rows", [])}
        else:
            print(f"[ERROR] 搜索失败: {data.get('msg')}")
            return {"total": 0, "rows": []}
    except Exception as e:
        print(f"[ERROR] 搜索请求失败: {e}")
        return {"total": 0, "rows": []}


def get_detail(bbbs: str) -> dict | None:
    """获取法律法规详情"""
    params = {"bbbs": bbbs}

    try:
        resp = requests.get(DETAIL_API, headers=HEADERS, params=params, timeout=30)
        resp.raise_for_status()
        data = resp.json()
        if data.get("code") == 200:
            return data.get("data")
        else:
            print(f"[WARN] 详情获取失败 ({bbbs}): {data.get('msg')}")
            return None
    except Exception as e:
        print(f"[ERROR] 详情请求失败 ({bbbs}): {e}")
        return None


def create_docx_from_html_content(title: str, content: str, output_path: Path) -> bool:
    """从 HTML 内容创建 Word 文档（处理富文本格式）"""
    try:
        doc = Document()

        # 添加标题
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()  # 空行

        # 处理内容（可能是 HTML 或纯文本）
        if content:
            # 清理 HTML 标签（简单处理）
            text = re.sub(r'<[^>]+>', '', content)  # 移除 HTML 标签
            text = re.sub(r'\s+', '\n', text)  # 规范化空白

            # 按段落分割
            paragraphs = [p.strip() for p in text.split('\n') if p.strip()]

            for para_text in paragraphs:
                # 检查是否是章节标题
                if re.match(r'^[第\s]*[一二三四五六七八九十百千万]+[章节条款]', para_text):
                    section_para = doc.add_paragraph()
                    section_run = section_para.add_run(para_text)
                    section_run.font.size = Pt(14)
                    section_run.font.bold = True
                else:
                    doc.add_paragraph(para_text)

        doc.save(output_path)
        return True
    except Exception as e:
        print(f"[ERROR] 创建文档失败 ({title}): {e}")
        return False


def sanitize_filename(name: str) -> str:
    """清理文件名"""
    # 移除不允许的字符
    name = re.sub(r'[<>:"/\\|?*]', '', name)
    # 限制长度
    if len(name) > 100:
        name = name[:100]
    return name.strip()


def crawl_category(category_name: str, category_ids: list[int], target_count: int,
                   downloaded_titles: set[str]) -> int:
    """爬取指定分类的法律"""
    print(f"\n{'='*60}")
    print(f"[INFO] 开始爬取 {category_name}...")
    print(f"[INFO] 分类代码: {category_ids}")
    print(f"{'='*60}")

    total_downloaded = 0
    page_num = 1
    page_size = 50

    while total_downloaded < target_count:
        # 搜索有效状态的法律
        result = search_laws(category_ids, page_num=page_num, page_size=page_size)
        total = result.get("total", 0)
        rows = result.get("rows", [])

        if not rows:
            print(f"[INFO] {category_name} 无更多数据")
            print(f"[INFO] 该分类共 {total} 条有效记录")
            break

        print(f"\n[INFO] {category_name} 第 {page_num} 页，获取 {len(rows)} 条（总 {total} 条）")

        for row in rows:
            if total_downloaded >= target_count:
                break

            bbbs = row.get("bbbs")
            title = row.get("title", "")

            if not bbbs or not title:
                continue

            # 检查重复
            if title in downloaded_titles:
                print(f"[SKIP] 已下载: {title[:50]}...")
                continue

            # 清理标题
            clean_title = sanitize_filename(title)
            output_path = OUTPUT_DIR / f"{clean_title}.docx"

            # 检查文件是否已存在
            if output_path.exists():
                print(f"[SKIP] 文件已存在: {clean_title[:50]}...")
                downloaded_titles.add(title)
                continue

            # 获取详情
            detail = get_detail(bbbs)
            if not detail:
                print(f"[FAIL] 获取详情失败: {title[:50]}...")
                continue

            # 获取内容
            content = detail.get("content") or ""
            content_len = len(content) if content else 0

            if content_len > 100:
                # 从内容创建文档
                if create_docx_from_html_content(title, content, output_path):
                    print(f"[OK] 创建成功: {clean_title[:50]}... ({content_len} 字)")
                    downloaded_titles.add(title)
                    total_downloaded += 1
                else:
                    print(f"[FAIL] 创建失败: {title[:50]}...")
            else:
                # 内容为空，尝试其他方式
                print(f"[WARN] 内容为空或太短: {title[:50]}... ({content_len} 字)")

                # 检查是否有 ossWordPath（尝试下载）
                oss_file = detail.get("ossFile", {})
                word_path = oss_file.get("ossWordPath")
                if word_path:
                    print(f"[INFO] 发现 Word 文件路径: {word_path}")
                    # 由于下载 API 可能有问题，暂时跳过

            # 延迟避免请求过快
            time.sleep(0.3)

        page_num += 1
        time.sleep(1)  # 页间延迟

    return total_downloaded


def main():
    """主函数"""
    print("=" * 60)
    print("国家法律法规数据库爬虫 - 修正版")
    print("=" * 60)
    print(f"[INFO] 目标目录: {OUTPUT_DIR}")
    print("[INFO] 目标数量: 约 1000 部有效法律文档")
    print("[INFO] 优先级: 宪法 > 法律 > 行政法规 > 司法解释 > 监察法规")
    print("[INFO] 只爬取有效状态（sxx=3）的文档")
    print("=" * 60)

    downloaded_titles: set[str] = set()
    total_count = 0

    # 按优先级爬取各分类（目标数量分配）
    categories_priority = [
        ("宪法", CATEGORY_CODES["宪法"], 50),           # 宪法类
        ("法律", CATEGORY_CODES["法律"], 350),         # 法律类
        ("行政法规", CATEGORY_CODES["行政法规"], 350), # 行政法规
        ("司法解释", CATEGORY_CODES["司法解释"], 250), # 司法解释
        ("监察法规", CATEGORY_CODES["监察法规"], 50),   # 监察法规
    ]

    for category_name, category_ids, target in categories_priority:
        count = crawl_category(category_name, category_ids, target, downloaded_titles)
        total_count += count
        print(f"\n[INFO] {category_name} 完成，本分类下载 {count} 部")

        if total_count >= 1000:
            print(f"\n[INFO] 已达到目标数量 {total_count}，停止爬取")
            break

    print("=" * 60)
    print(f"[DONE] 爬取完成！")
    print(f"[INFO] 共下载 {total_count} 部法律法规")
    print(f"[INFO] 文件保存在: {OUTPUT_DIR}")

    # 统计文件
    files = list(OUTPUT_DIR.glob("*.docx"))
    print(f"[INFO] 目录中共有 {len(files)} 个 .docx 文件")
    print("=" * 60)


if __name__ == "__main__":
    main()