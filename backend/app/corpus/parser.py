from __future__ import annotations

import hashlib
import re
from datetime import date
from pathlib import Path

from docx import Document

from app.domain.models import LawDocument, LegalArticle

_ARTICLE_PATTERN = re.compile(r"^(第[〇零一二三四五六七八九十百千万0-9]+条)(?:\s*)(.*)$")
_CONTEXT_PATTERN = re.compile(
    r"^第[〇零一二三四五六七八九十百千万0-9]+(?:编|章|节)\s*.*$"
)
_DATE_SUFFIX_PATTERN = re.compile(r"[-_](\d{4})(\d{2})(\d{2})$")


def parse_docx(path: str | Path) -> LawDocument:
    source = Path(path)
    paragraphs = [_normalize(item.text) for item in Document(source).paragraphs]
    paragraphs = [item for item in paragraphs if item]
    if not paragraphs:
        raise ValueError(f"{source.name} 没有可解析的正文")

    law_name, source_date = _metadata_from_filename(source)
    document_text = "\n".join(paragraphs)
    content_hash = hashlib.sha256(document_text.encode("utf-8")).hexdigest()
    version_key = source_date.isoformat() if source_date else content_hash[:12]
    document_id = _stable_id(law_name, version_key)
    articles = _split_articles(
        paragraphs=paragraphs,
        law_name=law_name,
        document_id=document_id,
        source_file=source.name,
        source_date=source_date,
    )
    return LawDocument(
        document_id=document_id,
        law_name=law_name,
        source_file=source.name,
        source_date=source_date,
        content_hash=content_hash,
        articles=articles,
    )


def _split_articles(
    *,
    paragraphs: list[str],
    law_name: str,
    document_id: str,
    source_file: str,
    source_date: date | None,
) -> list[LegalArticle]:
    result: list[LegalArticle] = []
    context: list[str] = []
    current_number: str | None = None
    current_content: list[str] = []
    current_context: list[str] = []

    def flush() -> None:
        if current_number is None:
            return
        content = "\n".join(current_content)
        result.append(
            LegalArticle(
                article_id=_stable_id(document_id, current_number),
                document_id=document_id,
                law_name=law_name,
                article_number=current_number,
                content=content,
                context_path=current_context,
                source_file=source_file,
                source_date=source_date,
            )
        )

    for paragraph in paragraphs:
        if _CONTEXT_PATTERN.match(paragraph):
            flush()
            current_number = None
            current_content = []
            context = [paragraph]
            continue
        match = _ARTICLE_PATTERN.match(paragraph)
        if match:
            flush()
            current_number = match.group(1)
            current_content = [paragraph]
            current_context = list(context)
        elif current_number is not None:
            current_content.append(paragraph)

    flush()
    if result:
        return result

    content = "\n".join(paragraphs)
    return [
        LegalArticle(
            article_id=_stable_id(document_id, "fallback-1"),
            document_id=document_id,
            law_name=law_name,
            article_number="全文",
            content=content,
            source_file=source_file,
            source_date=source_date,
            parse_mode="fallback",
        )
    ]


def _metadata_from_filename(path: Path) -> tuple[str, date | None]:
    stem = path.stem.strip()
    match = _DATE_SUFFIX_PATTERN.search(stem)
    if not match:
        return stem, None
    parsed_date = date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    return stem[: match.start()].strip("-_ "), parsed_date


def _normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u3000", " ")).strip()


def _stable_id(*parts: str) -> str:
    value = "|".join(_normalize(part).lower() for part in parts)
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:24]
